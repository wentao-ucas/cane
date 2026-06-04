# -*- coding: utf-8 -*-
"""
Markdown 转 Word (docx) 转换脚本
- 中文: 宋体 小四 (12pt)
- 英文/数字: Times New Roman
- 支持 **加粗** 解析
- 表格、图片嵌入
"""

import sys
import io

# Windows 编码处理
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

import re
import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime


# 字体设置
FONT_CN = '宋体'
FONT_EN = 'Times New Roman'
FONT_SIZE_NORMAL = Pt(12)  # 小四
FONT_SIZE_SMALL = Pt(10.5)  # 五号
FONT_SIZE_CAPTION = Pt(9)


def is_chinese_char(char):
    """判断是否是中文字符"""
    return '\u4e00' <= char <= '\u9fff' or char in '，。！？、；：""''（）【】《》—…'


def set_run_font(run, is_bold=False, is_italic=False, size=None):
    """设置run的字体，中文用宋体，英文用Times New Roman"""
    run.font.name = FONT_EN  # 西文字体
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CN)  # 中文字体
    run.font.size = size or FONT_SIZE_NORMAL
    run.bold = is_bold
    run.italic = is_italic


def add_formatted_text(paragraph, text, default_bold=False, default_italic=False, size=None):
    """
    添加格式化文本到段落，支持 **加粗** 和 *斜体* 解析
    """
    # 解析 **加粗** 和 *斜体* 标记
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|[^*`]+)'
    parts = re.findall(pattern, text)
    
    for part in parts:
        if not part:
            continue
        
        is_bold = default_bold
        is_italic = default_italic
        is_code = False
        content = part
        
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            is_bold = True
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            content = part[1:-1]
            is_italic = True
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            is_code = True
        
        if content:
            run = paragraph.add_run(content)
            if is_code:
                run.font.name = 'Consolas'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
                run.font.size = Pt(10)
            else:
                set_run_font(run, is_bold=is_bold, is_italic=is_italic, size=size)


def set_doc_styles(doc):
    """设置文档样式"""
    # Normal 样式
    style = doc.styles['Normal']
    style.font.name = FONT_EN
    style.font.size = FONT_SIZE_NORMAL
    style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CN)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
    
    # 标题样式
    heading_sizes = {1: 22, 2: 16, 3: 14, 4: 12}
    for i in range(1, 5):
        try:
            heading = doc.styles[f'Heading {i}']
            heading.font.name = FONT_EN
            heading.font.bold = True
            heading.font.size = Pt(heading_sizes.get(i, 12))
            heading._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CN)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(6)
        except:
            pass


def add_horizontal_line(doc):
    """添加水平分割线"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    # 使用底部边框模拟分割线
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)


def convert_md_to_docx(md_path, docx_path):
    """将Markdown文件转换为Word文档"""
    print(f"📄 转换: {md_path}")
    print(f"   → {docx_path}")
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    
    doc = Document()
    set_doc_styles(doc)
    
    i = 0
    in_code_block = False
    code_content = []
    in_table = False
    table_rows = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # ========== 代码块处理 ==========
        if stripped.startswith('```'):
            if in_code_block:
                # 代码块结束 - 创建表格样式的代码框
                if code_content:
                    # 用单单元格表格来显示代码块
                    code_table = doc.add_table(rows=1, cols=1)
                    code_table.style = 'Table Grid'
                    cell = code_table.cell(0, 0)
                    
                    # 设置背景色
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'F5F5F5')
                    cell._tc.get_or_add_tcPr().append(shading)
                    
                    cell.text = '\n'.join(code_content)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.name = 'Consolas'
                            run.font.size = Pt(9)
                    
                    doc.add_paragraph()  # 代码块后空行
                
                code_content = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
        
        # ========== 表格处理 ==========
        if '|' in stripped and stripped.startswith('|'):
            # 跳过分隔行
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                i += 1
                continue
            
            # 解析表格单元格
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            
            if not in_table:
                in_table = True
                table_rows = [cells]
            else:
                table_rows.append(cells)
            
            # 检查下一行是否还是表格
            next_is_table = False
            if i + 1 < len(lines):
                next_line = lines[i + 1].strip()
                if next_line.startswith('|') or re.match(r'^\|[\s\-:|]+\|$', next_line):
                    next_is_table = True
            
            if not next_is_table and table_rows:
                # 输出表格
                if len(table_rows) > 0 and len(table_rows[0]) > 0:
                    num_cols = max(len(row) for row in table_rows)
                    table = doc.add_table(rows=len(table_rows), cols=num_cols)
                    table.style = 'Table Grid'
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for row_idx, row_data in enumerate(table_rows):
                        for col_idx, cell_data in enumerate(row_data):
                            if col_idx < num_cols:
                                cell = table.cell(row_idx, col_idx)
                                # 清空默认段落
                                cell.paragraphs[0].clear()
                                # 添加格式化文本
                                para = cell.paragraphs[0]
                                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                is_header = (row_idx == 0)
                                add_formatted_text(para, cell_data, 
                                                  default_bold=is_header,
                                                  size=FONT_SIZE_SMALL)
                    
                    doc.add_paragraph()
                
                in_table = False
                table_rows = []
            
            i += 1
            continue
        
        # ========== 空行 ==========
        if not stripped:
            i += 1
            continue
        
        # ========== 标题 ==========
        if stripped.startswith('#'):
            match = re.match(r'^(#{1,6})\s+(.+)$', stripped)
            if match:
                level = len(match.group(1))
                text = match.group(2)
                # 移除标题中的格式标记用于显示
                clean_text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
                clean_text = re.sub(r'\*([^*]+)\*', r'\1', clean_text)
                heading = doc.add_heading(level=min(level, 4))
                add_formatted_text(heading, clean_text, default_bold=True,
                                  size=Pt({1: 22, 2: 16, 3: 14, 4: 12}.get(level, 12)))
                i += 1
                continue
        
        # ========== 引用块 ==========
        if stripped.startswith('>'):
            text = stripped[1:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.first_line_indent = Cm(0)
            add_formatted_text(p, text, default_italic=True, size=FONT_SIZE_SMALL)
            i += 1
            continue
        
        # ========== 分隔线 ==========
        if stripped in ['---', '***', '___']:
            add_horizontal_line(doc)
            i += 1
            continue
        
        # ========== 图片 ==========
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            
            # 处理相对路径
            base_dir = os.path.dirname(md_path)
            if img_path.startswith('../'):
                img_path = os.path.normpath(os.path.join(base_dir, img_path))
            elif not os.path.isabs(img_path):
                img_path = os.path.normpath(os.path.join(base_dir, img_path))
            
            if os.path.exists(img_path):
                try:
                    # 添加图片
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(5.5))
                    
                    # 添加图片说明
                    caption = doc.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_run = caption.add_run(f'图: {alt_text}')
                    set_run_font(cap_run, is_italic=True, size=FONT_SIZE_CAPTION)
                except Exception as e:
                    p = doc.add_paragraph(f'[图片加载失败: {alt_text}]')
            else:
                p = doc.add_paragraph(f'[图片不存在: {alt_text}]')
            
            i += 1
            continue
        
        # ========== 列表项 ==========
        list_match = re.match(r'^(\s*)[-*+]\s+(.+)$', line)
        if list_match:
            indent = len(list_match.group(1))
            text = list_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5 + indent * 0.3)
            bullet_run = p.add_run('• ')
            set_run_font(bullet_run)
            add_formatted_text(p, text)
            i += 1
            continue
        
        # 数字列表
        num_list_match = re.match(r'^(\s*)\d+\.\s+(.+)$', line)
        if num_list_match:
            indent = len(num_list_match.group(1))
            text = num_list_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5 + indent * 0.3)
            add_formatted_text(p, text)
            i += 1
            continue
        
        # ========== 普通段落 ==========
        # 移除链接语法但保留文本
        text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', stripped)
        p = doc.add_paragraph()
        add_formatted_text(p, text)
        
        i += 1
    
    # ========== 页脚 ==========
    add_horizontal_line(doc)
    footer = doc.add_paragraph()
    footer_text = f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M")} | 数据来源: akshare + USDA WASDE'
    add_formatted_text(footer, footer_text, size=FONT_SIZE_CAPTION, default_italic=True)
    
    # 保存
    doc.save(docx_path)
    print(f"   ✅ 保存成功!")


def main():
    """主函数"""
    print("="*60)
    print("📄 Markdown → Word 转换工具")
    print("   字体: 中文宋体小四, 英文Times New Roman")
    print("="*60)
    
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    articles_dir = os.path.join(base_dir, "articles")
    output_dir = os.path.join(base_dir, "output")
    
    for filename in os.listdir(articles_dir):
        if filename.endswith('.md'):
            md_path = os.path.join(articles_dir, filename)
            docx_name = filename.replace('.md', '.docx')
            docx_path = os.path.join(output_dir, docx_name)
            
            try:
                convert_md_to_docx(md_path, docx_path)
            except Exception as e:
                print(f"   ❌ 转换失败: {e}")
                import traceback
                traceback.print_exc()
    
    print("\n" + "="*60)
    print("✅ 全部转换完成!")
    print("="*60)


if __name__ == "__main__":
    main()
